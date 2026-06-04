from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import os

def create_ppt():
    prs = Presentation()
    prs.slide_width = Inches(16)
    prs.slide_height = Inches(9)
    
    # 自定义颜色
    BLUE_DARK = RGBColor(30, 64, 175)
    BLUE_LIGHT = RGBColor(96, 165, 250)
    GOLD = RGBColor(251, 191, 36)
    WHITE = RGBColor(255, 255, 255)
    GRAY = RGBColor(107, 114, 128)
    
    # ========== 封面页 ==========
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    title = slide.shapes.title
    title.text = "大都会人寿财富地图"
    title.text_frame.paragraphs[0].font.size = Pt(48)
    title.text_frame.paragraphs[0].font.color.rgb = WHITE
    title.text_frame.paragraphs[0].font.bold = True
    
    subtitle = slide.placeholders[1]
    subtitle.text = "微型纽约地标建筑模型 · 客户忠诚回馈项目"
    subtitle.text_frame.paragraphs[0].font.size = Pt(24)
    subtitle.text_frame.paragraphs[0].font.color.rgb = WHITE
    
    # 设置背景
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = BLUE_DARK
    
    # ========== 项目概览 ==========
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "项目概览"
    title.text_frame.paragraphs[0].font.size = Pt(36)
    title.text_frame.paragraphs[0].font.color.rgb = BLUE_DARK
    title.text_frame.paragraphs[0].font.bold = True
    
    content = slide.placeholders[1]
    content.text = "通过微型纽约地标建筑模型，为大都会人寿客户打造一个有温度、有故事、可参与、可收藏的忠诚度回馈系统。\n\n" \
                   "愿景：常伴左右，共驭美好未来\n" \
                   "Navigating life together"
    content.text_frame.paragraphs[0].font.size = Pt(20)
    
    # ========== 四大核心价值观 ==========
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    title = slide.shapes.title
    title.text = "四大核心价值观"
    title.text_frame.paragraphs[0].font.size = Pt(36)
    title.text_frame.paragraphs[0].font.color.rgb = BLUE_DARK
    title.text_frame.paragraphs[0].font.bold = True
    title.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    values = [
        {"icon": "🤝", "title": "携手共赢", "desc": "跨团队紧密协作，充分利用员工的多元化视角与团队力量"},
        {"icon": "🛡️", "title": "恪守正道", "desc": "践行我们的承诺，对自己和客户高度负责，保持诚信与公平"},
        {"icon": "⚡", "title": "质效先行", "desc": "拒绝无效忙碌，确保每一份时间与资源都精准投入"},
        {"icon": "🔭", "title": "远见卓识", "desc": "提前为未来做准备，洞察不同的可能性并主动适应变化"}
    ]
    
    for i, val in enumerate(values):
        x = Inches(1 + (i % 2) * 7)
        y = Inches(2.5 + (i // 2) * 3)
        width = Inches(6)
        height = Inches(2)
        
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(239, 246, 255)
        shape.line.color.rgb = RGBColor(147, 197, 253)
        shape.line.width = Pt(2)
        
        # 图标
        txBox = slide.shapes.add_textbox(x + Inches(0.5), y + Inches(0.2), Inches(1.5), Inches(1.5))
        tf = txBox.text_frame
        p = tf.add_paragraph()
        p.text = val["icon"]
        p.font.size = Pt(40)
        p.alignment = PP_ALIGN.CENTER
        
        # 标题
        txBox = slide.shapes.add_textbox(x + Inches(2.2), y + Inches(0.3), Inches(3), Inches(0.6))
        tf = txBox.text_frame
        p = tf.add_paragraph()
        p.text = val["title"]
        p.font.size = Pt(20)
        p.font.bold = True
        p.font.color.rgb = BLUE_DARK
        
        # 描述
        txBox = slide.shapes.add_textbox(x + Inches(2.2), y + Inches(1), Inches(3.5), Inches(0.8))
        tf = txBox.text_frame
        p = tf.add_paragraph()
        p.text = val["desc"]
        p.font.size = Pt(14)
        p.font.color.rgb = GRAY
    
    # ========== 五种设计方案 ==========
    designs = [
        {"id": "A", "title": "典雅金色典藏版", "style": "温暖金色 · 专业展示", 
         "desc": "透明长方形玻璃罩 + 金色底座，暖调夕阳纽约天际线背景"},
        {"id": "B", "title": "简约高级白金版", "style": "简洁白金 · 清新明亮",
         "desc": "极简白色+金色底座设计，清晰建筑标签，适合日常展示"},
        {"id": "C", "title": "复古穹顶收藏版", "style": "复古穹顶 · 怀旧典藏",
         "desc": "圆形玻璃穹顶设计，复古米色调，搭配老纽约背景"},
        {"id": "D", "title": "奢华黑金夜景版", "style": "深黑底座 · 金光璀璨",
         "desc": "黑色底座搭配金色灯光效果，呈现纽约夜景的奢华感"},
        {"id": "E", "title": "活力彩色水景版", "style": "多彩现代 · 生动水景",
         "desc": "彩色建筑（紫、金、蓝、绿）+ 水景基座，活泼且富有生命力"}
    ]
    
    for i in range(0, len(designs), 3):
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        
        if i == 0:
            title = slide.shapes.title
            title.text = "五种设计方案"
            title.text_frame.paragraphs[0].font.size = Pt(36)
            title.text_frame.paragraphs[0].font.color.rgb = BLUE_DARK
            title.text_frame.paragraphs[0].font.bold = True
            title.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        
        for j in range(min(3, len(designs) - i)):
            design = designs[i + j]
            x = Inches(0.5 + j * 5.2)
            y = Inches(2) if i == 0 else Inches(1.5)
            width = Inches(4.8)
            height = Inches(4.5 if i == 0 else 5)
            
            shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, width, height)
            shape.fill.solid()
            shape.fill.fore_color.rgb = WHITE
            shape.line.color.rgb = BLUE_LIGHT
            shape.line.width = Pt(2)
            
            # 方案标识
            txBox = slide.shapes.add_textbox(x + Inches(0.3), y + Inches(0.3), Inches(1), Inches(0.8))
            tf = txBox.text_frame
            p = tf.add_paragraph()
            p.text = f"方案{design['id']}"
            p.font.size = Pt(14)
            p.font.bold = True
            p.font.color.rgb = GOLD
            
            # 标题
            txBox = slide.shapes.add_textbox(x + Inches(0.3), y + Inches(1.1), Inches(4.2), Inches(0.6))
            tf = txBox.text_frame
            p = tf.add_paragraph()
            p.text = design["title"]
            p.font.size = Pt(18)
            p.font.bold = True
            p.font.color.rgb = BLUE_DARK
            
            # 风格
            txBox = slide.shapes.add_textbox(x + Inches(0.3), y + Inches(1.8), Inches(4.2), Inches(0.5))
            tf = txBox.text_frame
            p = tf.add_paragraph()
            p.text = design["style"]
            p.font.size = Pt(14)
            p.font.color.rgb = BLUE_LIGHT
            
            # 描述
            txBox = slide.shapes.add_textbox(x + Inches(0.3), y + Inches(2.4), Inches(4.2), Inches(1.8))
            tf = txBox.text_frame
            p = tf.add_paragraph()
            p.text = design["desc"]
            p.font.size = Pt(13)
            p.font.color.rgb = GRAY
            p.word_wrap = True
    
    # ========== 七大地标模型 ==========
    landmarks = [
        {"name": "大都会人寿大厦", "product": "都会守护百万医疗险", "desc": "基础防御"},
        {"name": "中央车站", "product": "都会无忧终身重疾险", "desc": "基础防御"},
        {"name": "克莱斯勒大厦", "product": "都会长福终身护理险", "desc": "品质进阶"},
        {"name": "范德比尔特一号大楼", "product": "都会长盈年金险", "desc": "品质进阶"},
        {"name": "洛克菲勒中心", "product": "都会颐年养老险", "desc": "财富规划"},
        {"name": "纽约公共图书馆总馆", "product": "都会盛世终身寿险", "desc": "财富规划"},
        {"name": "帝国大厦", "product": "都会臻传终身寿险", "desc": "财富规划"}
    ]
    
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    title = slide.shapes.title
    title.text = "七座纽约标志性地标"
    title.text_frame.paragraphs[0].font.size = Pt(36)
    title.text_frame.paragraphs[0].font.color.rgb = BLUE_DARK
    title.text_frame.paragraphs[0].font.bold = True
    title.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    for i, landmark in enumerate(landmarks):
        x = Inches(0.8 + (i % 4) * 3.8)
        y = Inches(2.5 + (i // 4) * 2.2)
        width = Inches(3.4)
        height = Inches(1.8)
        
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(248, 250, 252)
        shape.line.color.rgb = RGBColor(226, 232, 240)
        
        # 地标名称
        txBox = slide.shapes.add_textbox(x + Inches(0.3), y + Inches(0.2), Inches(3), Inches(0.5))
        tf = txBox.text_frame
        p = tf.add_paragraph()
        p.text = landmark["name"]
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = BLUE_DARK
        
        # 产品
        txBox = slide.shapes.add_textbox(x + Inches(0.3), y + Inches(0.8), Inches(3), Inches(0.4))
        tf = txBox.text_frame
        p = tf.add_paragraph()
        p.text = landmark["product"]
        p.font.size = Pt(11)
        p.font.color.rgb = GOLD
        
        # 描述标签
        tag_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x + Inches(0.3), y + Inches(1.3), Inches(1.5), Inches(0.3))
        tag_shape.fill.solid()
        tag_shape.fill.fore_color.rgb = BLUE_LIGHT
        tag_shape.line.color.rgb = BLUE_LIGHT
        
        txBox = slide.shapes.add_textbox(x + Inches(0.3), y + Inches(1.3), Inches(1.5), Inches(0.3))
        tf = txBox.text_frame
        p = tf.add_paragraph()
        p.text = landmark["desc"]
        p.font.size = Pt(10)
        p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.CENTER
    
    # ========== 逐级拼装流程 ==========
    steps = [
        "基础底板 - 财富之城规划底图",
        "大都会人寿大厦 - 都会守护百万医疗险",
        "中央车站 - 都会无忧终身重疾险",
        "克莱斯勒大厦 - 都会长福终身护理险",
        "范德比尔特一号大楼 - 都会长盈年金险",
        "洛克菲勒中心 - 都会颐年养老险",
        "纽约公共图书馆总馆 - 都会盛世终身寿险",
        "帝国大厦 + 玻璃罩 - 都会臻传终身寿险 · 完整拼装"
    ]
    
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    title = slide.shapes.title
    title.text = "逐级升级 · 拼筑财富之城"
    title.text_frame.paragraphs[0].font.size = Pt(36)
    title.text_frame.paragraphs[0].font.color.rgb = BLUE_DARK
    title.text_frame.paragraphs[0].font.bold = True
    title.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    for i, step in enumerate(steps):
        x = Inches(1)
        y = Inches(2 + i * 0.75)
        width = Inches(14)
        height = Inches(0.6)
        
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(239, 246, 255)
        shape.line.color.rgb = BLUE_LIGHT
        
        # 步骤编号
        txBox = slide.shapes.add_textbox(x + Inches(0.3), y + Inches(0.1), Inches(1), Inches(0.4))
        tf = txBox.text_frame
        p = tf.add_paragraph()
        p.text = str(i + 1)
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = BLUE_DARK
        p.alignment = PP_ALIGN.CENTER
        
        # 步骤内容
        txBox = slide.shapes.add_textbox(x + Inches(1.5), y + Inches(0.1), Inches(12), Inches(0.4))
        tf = txBox.text_frame
        p = tf.add_paragraph()
        p.text = step
        p.font.size = Pt(14)
        p.font.color.rgb = BLUE_DARK
    
    # ========== 项目声明 ==========
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    title = slide.shapes.title
    title.text = "项目声明"
    title.text_frame.paragraphs[0].font.size = Pt(36)
    title.text_frame.paragraphs[0].font.color.rgb = BLUE_DARK
    title.text_frame.paragraphs[0].font.bold = True
    title.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    content = [
        "本页面展示内容仅为设计方案展示用途，不构成最终产品交付标准。",
        "页面中所用图片均为AI生成，因技术局限性可能存在不可预期的误差。",
        "我们欢迎您的批评与指正，如有任何问题或建议，请联系：miaoweitech@aliyun.com",
        "",
        "© 2026 商业项目，版权归设计者所有 | 版本号：v1.0 | 提交日期：2026年6月4日"
    ]
    
    y = Inches(3)
    for line in content:
        txBox = slide.shapes.add_textbox(Inches(2), y, Inches(12), Inches(0.8))
        tf = txBox.text_frame
        p = tf.add_paragraph()
        p.text = line
        p.font.size = Pt(16 if line.startswith("©") else 14)
        p.font.color.rgb = GRAY if line else WHITE
        p.alignment = PP_ALIGN.CENTER
        y += Inches(0.7)
    
    # 保存PPT
    prs.save("大都会人寿财富地图_项目方案.pptx")
    print("PPT文件已生成：大都会人寿财富地图_项目方案.pptx")

if __name__ == "__main__":
    create_ppt()